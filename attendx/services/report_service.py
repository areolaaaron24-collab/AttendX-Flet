from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from attendx.core.config import Settings
from attendx.services.supabase_client import BackendError


class ReportService:
    def __init__(self, settings: Settings, teacher_service: Any, attendance_service: Any):
        self.settings = settings
        self.teacher_service = teacher_service
        self.attendance_service = attendance_service

    def subject_report_rows(self, schedule_id: str, attendance_date: str | None = None) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        schedules = self.teacher_service.my_schedules()
        schedule = next((s for s in schedules if str(s["id"]) == str(schedule_id)), None)
        if not schedule:
            raise BackendError("Schedule not found or not owned by this teacher.")
        enrollments = self.teacher_service.class_enrollments(schedule_id)
        records = self.attendance_service.teacher_records(schedule_id=schedule_id, attendance_date=attendance_date, limit=10_000)
        records_by_student = {r["student_id"]: r for r in records}
        rows = []
        for enrollment in enrollments:
            profile = enrollment.get("student", {})
            record = records_by_student.get(enrollment["student_id"], {})
            rows.append({
                "Account ID": profile.get("account_id", ""),
                "Student": " ".join(filter(None, [profile.get("first_name"), profile.get("middle_name"), profile.get("last_name"), profile.get("suffix")])),
                "Course": profile.get("course", ""),
                "Year": profile.get("year_level", ""),
                "Section": profile.get("section", ""),
                "Date": record.get("attendance_date", attendance_date or datetime.now(ZoneInfo("Asia/Manila")).date().isoformat()),
                "Time": record.get("attendance_time", ""),
                "Status": record.get("status", "NOT RECORDED"),
                "Method": record.get("method", ""),
            })
        return schedule, rows

    def _output_path(self, extension: str, subject_name: str) -> Path:
        folder = self.settings.exports_dir
        folder.mkdir(parents=True, exist_ok=True)
        safe = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in subject_name).strip("_") or "subject"
        stamp = datetime.now(ZoneInfo("Asia/Manila")).strftime("%Y%m%d_%H%M%S")
        return folder / f"AttendX_{safe}_{stamp}.{extension}"

    def export_excel(self, schedule_id: str, attendance_date: str | None = None) -> Path:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError as exc:
            raise BackendError("OpenPyXL is missing. Install requirements.txt to export Excel reports.") from exc
        schedule, rows = self.subject_report_rows(schedule_id, attendance_date)
        subject = schedule.get("subject", {}).get("subject_name", "Subject")
        path = self._output_path("xlsx", subject)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Attendance"
        sheet.append(["AttendX Smart Attendance System"])
        sheet.append(["Subject", subject])
        sheet.append(["Class", " ".join(filter(None, [schedule.get("class", {}).get("course"), schedule.get("class", {}).get("year_level"), schedule.get("class", {}).get("section")]))])
        sheet.append(["Schedule", f"{schedule.get('day', '')} {schedule.get('start_time', '')}–{schedule.get('end_time', '')}"])
        sheet.append([])
        headers = list(rows[0].keys()) if rows else ["Account ID", "Student", "Course", "Year", "Section", "Date", "Time", "Status", "Method"]
        sheet.append(headers)
        for item in rows:
            sheet.append([item.get(header, "") for header in headers])
        for cell in sheet[6]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="123F6B")
            cell.alignment = Alignment(wrap_text=True)
        sheet.freeze_panes = "A7"
        sheet.auto_filter.ref = f"A6:{chr(64 + len(headers))}{max(6, sheet.max_row)}"
        for column in sheet.columns:
            width = min(38, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
        workbook.save(path)
        return path

    def export_csv(self, schedule_id: str, attendance_date: str | None = None) -> Path:
        schedule, rows = self.subject_report_rows(schedule_id, attendance_date)
        path = self._output_path("csv", schedule.get("subject", {}).get("subject_name", "Subject"))
        headers = list(rows[0].keys()) if rows else ["Account ID", "Student", "Course", "Year", "Section", "Date", "Time", "Status", "Method"]
        with path.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=headers)
            writer.writeheader()
            writer.writerows(rows)
        return path

    def export_pdf(self, schedule_id: str, attendance_date: str | None = None) -> Path:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import landscape, A4
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        except ImportError as exc:
            raise BackendError("ReportLab is missing. Install requirements.txt to export PDF reports.") from exc
        schedule, rows = self.subject_report_rows(schedule_id, attendance_date)
        subject = schedule.get("subject", {}).get("subject_name", "Subject")
        path = self._output_path("pdf", subject)
        document = SimpleDocTemplate(str(path), pagesize=landscape(A4), rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)
        styles = getSampleStyleSheet()
        story = [Paragraph("AttendX Smart Attendance System", styles["Title"]),
                 Paragraph(f"{subject} — {schedule.get('day', '')} {schedule.get('start_time', '')}–{schedule.get('end_time', '')}", styles["Heading2"]),
                 Spacer(1, 12)]
        headers = list(rows[0].keys()) if rows else ["Account ID", "Student", "Course", "Year", "Section", "Date", "Time", "Status", "Method"]
        data = [headers] + [[str(row.get(h, ""))[:80] for h in headers] for row in rows]
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#123F6B")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CAD5E0")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F6FA")]),
        ]))
        story.append(table)
        document.build(story)
        return path
